from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "SRS_Fraud_Detection_System.docx"
MD_PATH = ROOT / "SRS_Fraud_Detection_System.md"
README_PATH = ROOT / "README.md"


TITLE = "Software Requirements Specification"
PROJECT = "Fraud Detection System for Financial Transactions"
AUTHOR = "Internship Capstone Project"
VERSION = "1.0"
DATE = "May 2026"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:{}".format(key)), str(value))


def style_table(table, header=True):
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": "D9E2EC"},
                bottom={"val": "single", "sz": "4", "color": "D9E2EC"},
                left={"val": "single", "sz": "4", "color": "D9E2EC"},
                right={"val": "single", "sz": "4", "color": "D9E2EC"},
            )
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        if header and row_index == 0:
            for cell in row.cells:
                set_cell_shading(cell, "1F4E79")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.bold = True


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    if p.runs:
        p.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)


def add_kv_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(4.8)
    hdr = table.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Details"
    for key, value in rows:
        row = table.add_row().cells
        row[0].text = key
        row[1].text = value
    style_table(table)
    doc.add_paragraph()
    return table


def add_matrix_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row_data in rows:
        row = table.add_row().cells
        for i, value in enumerate(row_data):
            row[i].text = value
    style_table(table)
    doc.add_paragraph()
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(31, 78, 121)
    return doc


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    run = p.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run(PROJECT)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(47, 84, 150)

    add_kv_table(
        doc,
        [
            ("Prepared for", "Internship Capstone Submission"),
            ("Prepared by", AUTHOR),
            ("Version", VERSION),
            ("Date", DATE),
            ("Document status", "Ready for review"),
        ],
    )

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    note.add_run(
        "This SRS defines the planned objectives, features, technical approach, "
        "constraints, and acceptance criteria for building a fraud detection platform."
    ).italic = True
    doc.add_page_break()


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = PROJECT
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(117, 117, 117)

        footer = section.footer.paragraphs[0]
        footer.text = "SRS Version 1.0"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(117, 117, 117)


def build_docx():
    doc = setup_document()
    add_header_footer(doc)
    add_cover(doc)

    add_heading(doc, "1. Introduction")
    add_heading(doc, "1.1 Purpose", level=2)
    add_body(
        doc,
        "The purpose of this Software Requirements Specification is to define the requirements "
        "for a Fraud Detection System for Financial Transactions. The system will analyze "
        "transaction data, identify suspicious patterns, assign fraud risk scores, and help "
        "financial teams review potentially fraudulent activity quickly."
    )
    add_heading(doc, "1.2 Project Scope", level=2)
    add_body(
        doc,
        "The project will deliver a machine learning based web application that accepts transaction "
        "records, evaluates them using a trained fraud detection model, stores prediction history, "
        "and presents alerts through an operational dashboard. The system is intended for capstone "
        "demonstration and academic review, with a design that can later be extended toward real "
        "banking or payment environments."
    )
    add_heading(doc, "1.3 Objectives", level=2)
    add_bullets(
        doc,
        [
            "Detect high-risk financial transactions using machine learning and rule-based validation.",
            "Provide transaction-level fraud probability, risk category, and reason indicators.",
            "Offer a dashboard for monitoring alerts, model metrics, and transaction trends.",
            "Maintain a clear project structure with backend APIs, frontend UI, ML pipeline, and documentation.",
            "Demonstrate secure handling of sensitive transaction and user information.",
        ],
    )
    add_heading(doc, "1.4 Intended Users", level=2)
    add_matrix_table(
        doc,
        ["User group", "Role in system", "Primary needs"],
        [
            ("Admin", "Manages system access and configuration", "User management, audit view, model status"),
            ("Fraud analyst", "Reviews flagged transactions", "Risk score, transaction context, decision support"),
            ("Bank/payment staff", "Uploads or enters transaction data", "Fast prediction and simple alert status"),
            ("Project evaluator", "Reviews internship submission", "Clear documentation, demo flow, measurable results"),
        ],
    )

    add_heading(doc, "2. Overall Description")
    add_heading(doc, "2.1 Product Perspective", level=2)
    add_body(
        doc,
        "The system will be designed as a standalone capstone application. It will combine a web "
        "frontend, REST API backend, relational database, and trained machine learning model. "
        "The model will be trained on historical transaction data and exposed through backend "
        "prediction endpoints."
    )
    add_heading(doc, "2.2 Product Functions", level=2)
    add_bullets(
        doc,
        [
            "User login and role-based access for admin and analyst users.",
            "Single transaction fraud prediction through a form-based interface.",
            "Batch transaction upload using CSV for multiple predictions.",
            "Real-time display of fraud risk score, predicted class, and alert priority.",
            "Dashboard charts for fraud count, approval count, amount distribution, and recent alerts.",
            "Transaction history with filter and search options.",
            "Model evaluation report including accuracy, precision, recall, F1-score, confusion matrix, and ROC-AUC.",
        ],
    )
    add_heading(doc, "2.3 Assumptions and Dependencies", level=2)
    add_bullets(
        doc,
        [
            "A public or synthetic financial transaction dataset will be used for training and testing.",
            "The project will use anonymized data only; no real customer banking data will be collected.",
            "The system will be developed within one week, so production banking integrations are out of scope.",
            "Internet access may be needed for dataset download, dependency installation, and deployment.",
        ],
    )

    add_heading(doc, "3. Functional Requirements")
    add_matrix_table(
        doc,
        ["ID", "Requirement", "Priority", "Acceptance criteria"],
        [
            ("FR-01", "The system shall allow authorized users to log in securely.", "High", "Invalid credentials are rejected and valid users reach the dashboard."),
            ("FR-02", "The system shall accept transaction details through a manual entry form.", "High", "A user can submit amount, merchant, location, time, card type, and account indicators."),
            ("FR-03", "The system shall support CSV upload for batch prediction.", "High", "A valid CSV returns predictions for all valid rows and reports row-level errors."),
            ("FR-04", "The system shall generate a fraud probability score for each transaction.", "High", "Each prediction includes probability, predicted class, and risk level."),
            ("FR-05", "The system shall classify results as Low, Medium, or High risk.", "High", "Risk labels are assigned using configurable thresholds."),
            ("FR-06", "The system shall save transaction predictions and timestamps.", "Medium", "Prediction history is visible after page refresh."),
            ("FR-07", "The system shall display recent high-risk alerts on the dashboard.", "High", "High-risk transactions appear in a clearly visible alerts table."),
            ("FR-08", "The system shall provide analytics charts for fraud trends.", "Medium", "Dashboard shows totals, fraud rate, amount distribution, and risk category count."),
            ("FR-09", "The system shall show model performance metrics.", "Medium", "Metrics page displays precision, recall, F1-score, ROC-AUC, and confusion matrix."),
            ("FR-10", "The system shall allow an analyst to mark a flagged transaction as reviewed.", "Medium", "Reviewed status and analyst note are stored and displayed."),
        ],
    )

    add_heading(doc, "4. Non-Functional Requirements")
    add_matrix_table(
        doc,
        ["Category", "Requirement"],
        [
            ("Performance", "Single transaction predictions should complete within 2 seconds under normal demo conditions."),
            ("Usability", "The dashboard should be simple enough for a reviewer to understand within 2 minutes."),
            ("Reliability", "The system should handle invalid form fields and malformed CSV rows without crashing."),
            ("Security", "Passwords must be hashed, sensitive fields must not be logged in plain text, and protected routes must require authentication."),
            ("Privacy", "Only anonymized or synthetic transaction data will be used in the capstone."),
            ("Maintainability", "Frontend, backend, ML training, and documentation should be organized in separate folders."),
            ("Portability", "The project should run locally using documented setup commands and optionally be deployable to a cloud platform."),
        ],
    )

    add_heading(doc, "5. System Features")
    add_heading(doc, "5.1 Fraud Prediction Engine", level=2)
    add_body(
        doc,
        "The fraud prediction engine will preprocess input transaction features, apply the trained "
        "model, and return a fraud probability. Candidate algorithms include Logistic Regression, "
        "Random Forest, XGBoost, and Isolation Forest. The final model will be selected based on "
        "recall, precision, F1-score, and explainability."
    )
    add_heading(doc, "5.2 Analyst Dashboard", level=2)
    add_body(
        doc,
        "The dashboard will summarize transaction volume, fraud rate, top risk categories, recent "
        "alerts, and review status. It will prioritize clarity over decoration because the target "
        "workflow is operational monitoring."
    )
    add_heading(doc, "5.3 Batch Upload and History", level=2)
    add_body(
        doc,
        "Users will upload CSV files containing multiple transactions. The backend will validate "
        "required columns, run predictions, store results, and return an exportable results table."
    )

    add_heading(doc, "6. Technical Requirements")
    add_matrix_table(
        doc,
        ["Layer", "Recommended technology", "Purpose"],
        [
            ("Frontend", "React.js or plain HTML/CSS/JavaScript", "Dashboard, forms, charts, login screens"),
            ("Backend", "Python Flask or FastAPI", "REST APIs, authentication, prediction endpoints"),
            ("Machine Learning", "Python, pandas, scikit-learn, imbalanced-learn", "Training, preprocessing, evaluation, model persistence"),
            ("Database", "SQLite for demo; PostgreSQL for extension", "Users, transactions, prediction history, analyst reviews"),
            ("Visualization", "Chart.js, Recharts, or Plotly", "Fraud trends and model metrics"),
            ("Deployment", "Render, Railway, Vercel plus backend host, or local demo", "Live demo or reviewer-accessible build"),
        ],
    )

    add_heading(doc, "7. Data Requirements")
    add_body(
        doc,
        "The dataset should contain transaction-level records with numeric, categorical, and temporal "
        "features. Because fraud datasets are usually imbalanced, the training process must address "
        "class imbalance using techniques such as stratified splitting, class weights, SMOTE, or "
        "threshold tuning."
    )
    add_matrix_table(
        doc,
        ["Data field", "Description", "Example"],
        [
            ("transaction_amount", "Transaction value", "2499.00"),
            ("transaction_time", "Date/time or derived time feature", "2026-05-10 14:20"),
            ("merchant_category", "Merchant type or category", "Electronics"),
            ("location", "Transaction location or region", "Bengaluru"),
            ("payment_method", "Card, UPI, wallet, net banking", "Card"),
            ("account_age_days", "Customer account age", "420"),
            ("previous_failed_attempts", "Recent failed transaction count", "2"),
            ("is_fraud", "Training label", "0 or 1"),
        ],
    )

    add_heading(doc, "8. External Interface Requirements")
    add_heading(doc, "8.1 User Interface", level=2)
    add_bullets(
        doc,
        [
            "Login page for authorized access.",
            "Dashboard page with key metrics and recent alerts.",
            "Predict transaction page for manual input.",
            "Batch upload page for CSV prediction.",
            "Transaction history page with filters for risk level, amount, and status.",
            "Model performance page for evaluation metrics.",
        ],
    )
    add_heading(doc, "8.2 API Endpoints", level=2)
    add_matrix_table(
        doc,
        ["Endpoint", "Method", "Purpose"],
        [
            ("/api/auth/login", "POST", "Authenticate user and return session/token."),
            ("/api/predict", "POST", "Predict fraud risk for one transaction."),
            ("/api/predict/batch", "POST", "Upload CSV and predict multiple transactions."),
            ("/api/transactions", "GET", "List transaction prediction history."),
            ("/api/transactions/{id}/review", "PATCH", "Update analyst review status and notes."),
            ("/api/model/metrics", "GET", "Return model evaluation metrics."),
        ],
    )

    add_heading(doc, "9. Architecture and Design Approach")
    add_body(
        doc,
        "The planned architecture follows a modular design. The frontend sends transaction data to "
        "the backend API. The backend validates data, loads the trained model, performs prediction, "
        "stores the result, and returns a response to the UI. The ML training pipeline is kept "
        "separate from runtime prediction so that model training and application serving remain clean."
    )
    add_matrix_table(
        doc,
        ["Module", "Responsibility"],
        [
            ("Frontend UI", "Collects user inputs, displays predictions, charts, and transaction tables."),
            ("API service", "Handles authentication, validation, prediction requests, and data access."),
            ("ML pipeline", "Cleans data, engineers features, trains model, evaluates performance, and saves model artifact."),
            ("Database", "Stores users, transactions, prediction outputs, review status, and audit timestamps."),
            ("Documentation", "Explains setup, usage, dataset, model results, and project limitations."),
        ],
    )

    add_heading(doc, "10. Database Design")
    add_matrix_table(
        doc,
        ["Table", "Key fields", "Purpose"],
        [
            ("users", "id, name, email, password_hash, role, created_at", "Stores authorized users."),
            ("transactions", "id, amount, merchant_category, location, payment_method, created_at", "Stores submitted transactions."),
            ("predictions", "id, transaction_id, fraud_probability, risk_level, predicted_label, model_version", "Stores model outputs."),
            ("reviews", "id, transaction_id, analyst_id, status, note, reviewed_at", "Stores analyst decisions and comments."),
        ],
    )

    add_heading(doc, "11. Security and Privacy")
    add_bullets(
        doc,
        [
            "Use password hashing for all stored user credentials.",
            "Restrict dashboard and prediction routes to authenticated users.",
            "Validate and sanitize API inputs to reduce injection and malformed data risks.",
            "Avoid storing real card numbers, CVV values, bank account numbers, or personally identifiable information.",
            "Use anonymized sample data for screenshots, demos, and LinkedIn posts.",
        ],
    )

    add_heading(doc, "12. Testing Strategy")
    add_matrix_table(
        doc,
        ["Test type", "Planned checks"],
        [
            ("Unit testing", "Validate preprocessing functions, risk threshold logic, and API validators."),
            ("API testing", "Check login, prediction, batch upload, history, and review endpoints."),
            ("Model testing", "Evaluate confusion matrix, precision, recall, F1-score, ROC-AUC, and false negative rate."),
            ("UI testing", "Confirm dashboard navigation, form validation, chart display, and responsive layout."),
            ("Integration testing", "Verify frontend-to-backend prediction flow and database persistence."),
        ],
    )

    add_heading(doc, "13. Project Timeline")
    add_matrix_table(
        doc,
        ["Day", "Work planned", "Deliverable"],
        [
            ("Day 1", "Finalize requirements, dataset selection, and SRS", "Approved SRS and project plan"),
            ("Day 2", "Data cleaning, feature engineering, and exploratory analysis", "Prepared training dataset"),
            ("Day 3", "Train baseline and improved ML models", "Saved model and metrics report"),
            ("Day 4", "Build backend APIs and database schema", "Working prediction API"),
            ("Day 5", "Build frontend dashboard and upload forms", "Usable web interface"),
            ("Day 6", "Integrate frontend, backend, and ML model", "End-to-end demo"),
            ("Day 7", "Testing, README, screenshots, and final submission", "Complete capstone package"),
        ],
    )

    add_heading(doc, "14. Risks and Mitigation")
    add_matrix_table(
        doc,
        ["Risk", "Impact", "Mitigation"],
        [
            ("Highly imbalanced fraud data", "Model may miss fraud cases", "Use recall-focused evaluation, class weights, and threshold tuning."),
            ("Dataset quality issues", "Poor model accuracy", "Clean missing values, remove duplicates, document limitations."),
            ("Too much scope for one week", "Incomplete project", "Prioritize prediction API, dashboard, README, and model metrics first."),
            ("Deployment difficulty", "No live demo", "Prepare local demo instructions and screenshots as fallback."),
            ("Sensitive data concerns", "Privacy risk", "Use anonymized or synthetic data only."),
        ],
    )

    add_heading(doc, "15. Acceptance Criteria")
    add_numbered(
        doc,
        [
            "The SRS is complete and clearly explains objectives, features, requirements, design approach, and timeline.",
            "The project repository contains source code for ML training, backend APIs, frontend UI, and documentation.",
            "A user can run the application locally using the README instructions.",
            "A transaction can be submitted and a fraud prediction result is returned.",
            "Batch CSV prediction works for multiple transactions.",
            "The dashboard displays recent alerts and summary metrics.",
            "The model evaluation report includes precision, recall, F1-score, ROC-AUC, and confusion matrix.",
            "The final submission includes the SRS document, project description, README, and demo link or screenshots.",
        ],
    )

    add_heading(doc, "16. Glossary")
    add_matrix_table(
        doc,
        ["Term", "Meaning"],
        [
            ("Fraud probability", "A numeric score representing the model's estimated chance that a transaction is fraudulent."),
            ("False negative", "A fraudulent transaction incorrectly predicted as legitimate."),
            ("Recall", "The percentage of actual fraud cases correctly detected."),
            ("Precision", "The percentage of predicted fraud cases that are actually fraud."),
            ("ROC-AUC", "A model evaluation metric that measures class separation ability."),
            ("Risk threshold", "A configured probability value used to classify Low, Medium, or High risk."),
        ],
    )

    doc.save(DOCX_PATH)


def build_markdown():
    md = f"""# {TITLE}: {PROJECT}

**Version:** {VERSION}  
**Date:** {DATE}  
**Prepared for:** Internship Capstone Submission

## Project Overview

This project is a machine learning based fraud detection system for financial transactions. It will analyze transaction features, predict whether a transaction is suspicious, assign a fraud probability score, and display fraud alerts through a dashboard.

## Objectives

- Detect high-risk financial transactions using machine learning.
- Provide transaction-level fraud probability and risk category.
- Support manual transaction prediction and batch CSV upload.
- Display fraud trends, alerts, and model performance metrics.
- Use anonymized or synthetic data only.

## Core Features

- Secure login for admin and fraud analyst users.
- Manual transaction prediction form.
- CSV batch upload and prediction results.
- Fraud probability, predicted label, and Low/Medium/High risk classification.
- Dashboard with fraud rate, recent alerts, amount trends, and model metrics.
- Transaction history with review status and analyst notes.

## Suggested Tech Stack

- Frontend: React.js
- Backend: Python Flask or FastAPI
- Machine Learning: pandas, scikit-learn, imbalanced-learn
- Database: SQLite for demo, PostgreSQL for extension
- Visualization: Chart.js, Recharts, or Plotly

## One-Week Development Plan

| Day | Work planned | Deliverable |
| --- | --- | --- |
| Day 1 | Finalize requirements and SRS | Approved SRS |
| Day 2 | Dataset preparation and EDA | Clean training dataset |
| Day 3 | Train and evaluate models | Saved model and metrics |
| Day 4 | Build backend APIs | Working prediction API |
| Day 5 | Build frontend dashboard | Usable web UI |
| Day 6 | Integrate full system | End-to-end demo |
| Day 7 | Testing and documentation | Final submission package |

## Submission Checklist

- SRS document
- Project source code
- README with setup instructions
- Project description
- Screenshots or live demo link
- LinkedIn post sharing the SRS/project
"""
    MD_PATH.write_text(md, encoding="utf-8")


def build_readme():
    readme = f"""# {PROJECT}

This folder contains the Software Requirements Specification for the internship capstone project: **{PROJECT}**.

## Files

- `SRS_Fraud_Detection_System.docx` - polished SRS document for submission.
- `SRS_Fraud_Detection_System.md` - lightweight Markdown version of the SRS summary.

## Short Project Description

The Fraud Detection System for Financial Transactions is a machine learning based web application that predicts whether a transaction is suspicious. It will provide fraud probability, risk classification, batch CSV prediction, transaction history, analyst review status, and a dashboard for monitoring fraud trends.

## LinkedIn Post Draft

I am working on my internship capstone project: **Fraud Detection System for Financial Transactions**.

As the first step, I prepared a Software Requirements Specification that defines the project scope, objectives, system features, technical requirements, machine learning approach, database design, testing plan, risks, and one-week development timeline.

The planned system will use transaction data to predict fraud risk, classify alerts as Low/Medium/High risk, and provide a dashboard for analysts to review suspicious transactions.

#Internship #CapstoneProject #MachineLearning #FraudDetection #SoftwareRequirements #DataScience
"""
    README_PATH.write_text(readme, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_markdown()
    build_readme()
    print(DOCX_PATH)
    print(MD_PATH)
    print(README_PATH)
